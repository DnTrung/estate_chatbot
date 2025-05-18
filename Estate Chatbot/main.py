import pandas as pd
import uuid
import tkinter as tk
from tkinter import ttk, messagebox
import threading
import logging
import random
from datetime import datetime
import json
from tkinter import filedialog
import os
from docx import Document as DocxDocument
import PyPDF2

# Import our custom modules
from chatbot import RealEstateChatbot
from user_context_db import UserContextDatabase

# Set up logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler("chatbot.log"),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

class RealEstateApp:
    def __init__(self, master):
        self.master = master
        self.master.title("Real Estate AI Assistant")
        self.master.geometry("1400x900")
        
        # Load property data
        self.load_data()
        
        # Initialize database
        self.db = UserContextDatabase()
        
        self.document = None
        # Initialize chatbot
        self.chatbot = RealEstateChatbot(self.property_data,self.document)
        
        # Generate a unique conversation ID
        self.conversation_id = str(uuid.uuid4())
        
        # Create temporary user ID for demo
        self.user_id = self.chatbot.user_id
        self.db.add_user(self.user_id)
        self.db.create_conversation(self.conversation_id, self.user_id)
        
        # Create UI frames
        self.create_frames()
        
        # Add welcome message
        self.add_bot_message("Xin chào! Tôi là trợ lý AI chuyên về bất động sản. Tôi có thể giúp bạn tìm kiếm căn hộ, nhà phố hoặc biệt thự theo nhu cầu của bạn. Bạn đang tìm kiếm bất động sản như thế nào? Hoặc tôi có thể giúp bạn giải đáp các thắc mắc liên quan đến bất động sản")
        
        # Start the periodic updates
        self.update_user_info()
    
    def load_data(self):
        """Load and prepare property data"""
        try:
            self.property_data = pd.read_csv("data/vietnam_housing_dataset.csv")
            
            # Generate descriptions
            def generate_description(row):
                return (
                    f"Căn hộ tại {row['Address']}, diện tích {row['Area']}m², "
                    f"{row['Bedrooms']} phòng ngủ, {row['Bathrooms']} phòng tắm, "
                    f"hướng {row['House direction']}, ban công hướng {row['Balcony direction']}. "
                    f"Nội thất: {row['Furniture state']}. "
                    f"Pháp lý: {row['Legal status']}. "
                    f"Mức giá: {row['Price']} tỷ VNĐ."
                )
            
            self.property_data["description"] = self.property_data.apply(generate_description, axis=1)
            
            # Clean data
            self.property_data = self.property_data.fillna({
                'Bedrooms': 0, 
                'Bathrooms': 0,
                'House direction': 'Không xác định',
                'Balcony direction': 'Không xác định',
                'Furniture state': 'Không xác định',
                'Legal status': 'Không xác định'
            })
            
            logger.info(f"Loaded {len(self.property_data)} properties")
        
        except Exception as e:
            logger.error(f"Error loading data: {e}")
            messagebox.showerror("Error", f"Failed to load property data: {e}")
            self.property_data = pd.DataFrame()  # Empty dataframe as fallback
    
    def create_frames(self):
        """Create UI frames and components"""
        # Main container
        container = ttk.PanedWindow(self.master, orient=tk.HORIZONTAL)
        container.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Chat section (left side)
        chat_frame = ttk.LabelFrame(container, text="AI Assistant")
        container.add(chat_frame, weight=3)
        
        # Chat display
        self.chat_display = tk.Text(chat_frame, wrap=tk.WORD, state='disabled')
        self.chat_display.tag_configure("user", foreground="blue")
        self.chat_display.tag_configure("bot", foreground="green")
        self.chat_display.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Chat input
        input_frame = ttk.Frame(chat_frame)
        input_frame.pack(fill=tk.X, padx=5, pady=5)
        
        self.chat_input = ttk.Entry(input_frame)
        self.chat_input.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))
        self.chat_input.bind("<Return>", self.send_message)
        
        send_button = ttk.Button(input_frame, text="Send", command=self.send_message)
        send_button.pack(side=tk.RIGHT)
        
        # Control panel (right side)
        control_panel = ttk.Frame(container)
        container.add(control_panel, weight=1)
        
        # User profile section
        user_frame = ttk.LabelFrame(control_panel, text="User Profile")
        user_frame.pack(fill=tk.X, expand=False, padx=5, pady=5)
        
        # User info fields
        fields = [
            ("Name", "name"), 
            ("Age", "age"), 
            ("Gender", "gender"),
            ("Income Level", "income_level"),
            ("Budget", "budget"),
            ("Family Info", "family_info"),
            ("Hobbies", "hobbies")
            
        ]
        
        self.user_vars = {}
        
        for label_text, field_name in fields:
            frame = ttk.Frame(user_frame)
            frame.pack(fill=tk.X, padx=5, pady=2)
            
            label = ttk.Label(frame, text=f"{label_text}:")
            label.pack(side=tk.LEFT)
            
            var = tk.StringVar()
            self.user_vars[field_name] = var
            
            entry = ttk.Entry(frame, textvariable=var)
            entry.pack(side=tk.RIGHT, fill=tk.X, expand=True)
        
        # Save profile button
        save_button = ttk.Button(user_frame, text="Save Profile", command=self.save_user_profile)
        save_button.pack(fill=tk.X, padx=5, pady=5)
        
        # Staff suggestion section
        staff_frame = ttk.LabelFrame(control_panel, text="Staff Suggestions")
        staff_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.suggestion_text = tk.Text(staff_frame, height=10)
        self.suggestion_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        suggestion_button = ttk.Button(staff_frame, text="Send Suggestion", command=self.send_message)
        suggestion_button.pack(fill=tk.X, padx=5, pady=5)

        # Add file upload button
        upload_button = ttk.Button(staff_frame, text="Upload File", command=self.upload_file)
        upload_button.pack(fill=tk.X, padx=5, pady=5)
        
        # Quick suggestions
        quick_frame = ttk.LabelFrame(control_panel, text="Quick Suggestions")
        quick_frame.pack(fill=tk.X, expand=False, padx=5, pady=5)
        
        suggestions = [
            "Đề xuất căn hộ có nội thất đầy đủ",
            "Nhấn mạnh về tính pháp lý đầy đủ"
        ]
        
        for suggestion in suggestions:
            button = ttk.Button(
                quick_frame, 
                text=suggestion, 
                command=lambda s=suggestion: self.use_quick_suggestion(s)
            )
            button.pack(fill=tk.X, padx=5, pady=2)
        
        # User context display
        context_frame = ttk.LabelFrame(control_panel, text="User Context & Preferences")
        context_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        self.context_text = tk.Text(context_frame, height=10, state='disabled')
        self.context_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        # Add clear button
        clear_button = ttk.Button(context_frame, text="Clear All", command=self.clear_all)
        clear_button.pack(fill=tk.X, padx=5, pady=5)


    def upload_file(self):
        """Allow staff to upload a file"""
        file_path = filedialog.askopenfilename(
            title="Select a File",
            filetypes=(
                ("PDF files", "*.pdf"),
                ("Word documents", "*.doc;*.docx"),
                ("Text files", "*.txt"),
                ("All files", "*.*"),
            ),
        )
        if file_path:
            file_name = os.path.basename(file_path)
            self.suggestion_text.insert(tk.END, f"Uploaded File: {file_name}\nPath: {file_path}\n")
            self.add_staff_suggestion(f"Uploaded File: {file_name}")
            
            # Read content based on file extension
            ext = os.path.splitext(file_path)[1].lower()
            document = ""
            try:
                if ext == ".txt":
                    with open(file_path, "r", encoding="utf-8") as file:
                        document = file.read()
                elif ext == ".pdf":
                    with open(file_path, "rb") as file:
                        reader = PyPDF2.PdfReader(file)
                        document = "\n".join(page.extract_text() or "" for page in reader.pages)
                elif ext in [".doc", ".docx"]:
                    doc = DocxDocument(file_path)
                    document = "\n".join([para.text for para in doc.paragraphs])
                else:
                    document = "Unsupported file type."
            except Exception as e:
                document = f"Error reading file: {e}"
            self.document = document

    def send_message(self, event=None):
        """Send user message to the chatbot"""
        if self.chat_input.get():
            message = self.chat_input.get()
        else:
            message = self.suggestion_text.get("1.0", tk.END).strip()
        
        # Clear input field
        self.chat_input.delete(0, tk.END)
        
        self.add_user_message(message)
        
        # Save message to database
        self.db.add_message(self.conversation_id, "user", message)
        
        # Get staff suggestion if any
        staff_suggestion = self.suggestion_text.get("1.0", tk.END).strip()
        if staff_suggestion:
            self.suggestion_text.delete("1.0", tk.END)
        else:
            staff_suggestion = None
        if staff_suggestion:
            self.add_staff_suggestion(staff_suggestion)
        
        # Process message and get response
        response = self.chatbot.process_message(message, staff_suggestion)
        
        # Add response to chat display
        
            
        self.add_bot_message(response)
        
        # Save response to database
        message_id = self.db.add_message(self.conversation_id, "bot", response)
        
        # If there was a staff suggestion, save it
        if staff_suggestion:
            self.db.add_staff_suggestion(self.conversation_id, message_id, staff_suggestion)
        
        # Update user preferences based on current chatbot state
        self.update_preferences_from_chatbot()
    
    def add_user_message(self, message):
        """Add user message to chat display"""
        staff_suggestion = self.suggestion_text.get("1.0", tk.END).strip()
        self.chat_display.config(state='normal')
        if staff_suggestion in message: 
            message.replace(staff_suggestion,"")
            self.chat_display.insert(tk.END, f"You: {message}\n\n", "user")
        self.chat_display.config(state='disabled')
        self.chat_display.see(tk.END)
    
    def add_bot_message(self, message):
        """Add bot message to chat display"""
        self.chat_display.config(state='normal')
        self.chat_display.insert(tk.END, f"Assistant: {message}\n\n", "bot")
        self.chat_display.config(state='disabled')
        self.chat_display.see(tk.END)
    
    def add_staff_suggestion(self, message):
        """Add staff suggestion to chat display"""
        self.chat_display.config(state='normal')
        self.chat_display.insert(tk.END, f"Staff Suggestion: {message}\n\n", "staff")
        self.chat_display.config(state='disabled')
        self.chat_display.see(tk.END)
    
    def save_user_profile(self):
        """Save user profile information"""
        profile = {}
        for field, var in self.user_vars.items():
            value = var.get()
            if field == 'age' and value:
                try:
                    value = int(value)
                except ValueError:
                    messagebox.showerror("Error", "Age must be a number")
                    return
            
            profile[field] = value
        
        # Update user in database
        self.db.update_user(self.user_id, **profile)
        
        messagebox.showinfo("Success", "User profile saved successfully")
    
    def send_staff_suggestion(self):
        """Prepare staff suggestion to be used with next message"""
        # The suggestion will be used when the next user message is sent
        pass
    
    def use_quick_suggestion(self, suggestion):
        """Use a quick suggestion template"""
        self.suggestion_text.delete("1.0", tk.END)
        self.suggestion_text.insert("1.0", suggestion)
    
    def update_preferences_from_chatbot(self):
        """Update user preferences based on chatbot state"""
        preferences = self.chatbot.get_user_preferences()
        
        # Convert to database format
        db_preferences ={
            'min_price': preferences['min_price'],
            'max_price': preferences['max_price'],
            'min_area': preferences['min_area'],
            'max_area': preferences['max_area'],
            'min_bedrooms': preferences['bedrooms'],
            'min_bathrooms': preferences['bathrooms'],
            'preferred_districts': preferences['locations'],
            'preferred_direction': preferences['house_direction'],
            'furniture_state': preferences['furniture_state'],
            'legal_state': preferences['legal_state']     
        }
        
        # Update in database
        self.db.update_user_preferences(self.user_id, **db_preferences)
    
    def update_user_info(self):
        """Update user info display"""
        # Get user from database
        user = self.db.get_user(self.user_id)

        # Get user preferences
        preferences = self.db.get_user_preferences(self.user_id)

        # Update UI
        self.context_text.config(state='normal')
        self.context_text.delete("1.0", tk.END)

        if user:
            self.context_text.insert(tk.END, "USER PROFILE:\n")
            for key, value in user.items():
                if key not in ['user_id', 'created_at', 'updated_at'] and value:
                    self.context_text.insert(tk.END, f"{key}: {value}\n")

        if preferences:
            self.context_text.insert(tk.END, "\nREAL ESTATE PREFERENCES:\n")
            for key, value in preferences.items():
                if key == 'preferred_districts' and value:
                    # Decode JSON-encoded string if necessary
                    if isinstance(value, str):
                        try:
                            value = json.loads(value)  # Decode JSON string to a Python list
                        except json.JSONDecodeError:
                            pass  # If decoding fails, keep the original value
                    # Display preferred districts as a comma-separated list
                    districts = ", ".join(value) if isinstance(value, list) else value
                    self.context_text.insert(tk.END, f"{key}: {districts}\n")
                elif key not in ['preference_id', 'user_id'] and value:
                    self.context_text.insert(tk.END, f"{key}: {value}\n")

        self.context_text.config(state='disabled')

    # Schedule next update
        self.master.after(1000, self.update_user_info)
    
    def clear_all(self):
        """Clear all preferences, chat history, and reset the chatbot"""
        # Clear chat display
        self.chat_display.config(state='normal')
        self.chat_display.delete('1.0', tk.END)
        self.chat_display.config(state='disabled')
        
        # Clear staff suggestion
        self.suggestion_text.delete('1.0', tk.END)
        
        # Reset chatbotf
        self.chatbot.reset_conversation()
        
        # Reset user preferences in database
        self.db.update_user_preferences(self.user_id, 
            min_price=None,
            max_price=None,
            min_area=None,
            max_area=None,
            min_bedrooms=None,
            min_bathrooms=None,
            preferred_districts=[],
            preferred_direction=None,
            legal_state=None,
            furniture_state=None,
        )
        
        # Update UI
        self.update_user_info()
        
        # Add welcome message
        self.add_bot_message("Xin chào! Tôi là trợ lý AI chuyên về bất động sản. Tôi có thể giúp bạn tìm kiếm căn hộ, nhà phố hoặc biệt thự theo nhu cầu của bạn. Bạn đang tìm kiếm bất động sản như thế nào?")

def main():
    # Set up main application window
    root = tk.Tk()
    app = RealEstateApp(root)
    
    # Run the application
    root.mainloop()

if __name__ == "__main__":
    main()